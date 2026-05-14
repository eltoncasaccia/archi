import html as _html
import io
import logging
import re
from datetime import datetime, timezone

from docx import Document
from weasyprint import HTML

from app.services.supabase_client import get_supabase

logger = logging.getLogger(__name__)

BUCKET = "proposals"


def _build_context(proposal: dict, session: dict) -> dict:
    price = proposal.get("total_price")
    return {
        "{{client_name}}": proposal.get("client_name") or session.get("client_name") or "—",
        "{{client_email}}": proposal.get("client_email") or session.get("client_email") or "—",
        "{{client_document}}": session.get("client_document") or "—",
        "{{scope_description}}": session.get("discovery_summary") or "—",
        "{{phases_description}}": session.get("phases_plan") or "—",
        "{{total_price}}": f"R$ {price:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if price else "A definir",
        "{{total_days}}": str(proposal.get("total_days") or "A definir"),
        "{{admin_notes}}": proposal.get("admin_notes") or "—",
    }


# ---------------------------------------------------------------------------
# DOCX generation
# ---------------------------------------------------------------------------

    for para in doc.paragraphs:
        replace_in_para(para)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    replace_in_para(para)


def _generate_docx_bytes(proposal: dict, session: dict) -> bytes:
    from docx.shared import Pt, RGBColor, Cm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()

    for s in doc.sections:
        s.page_height = Cm(29.7)
        s.page_width = Cm(21.0)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.2)

    # Background da página: mesma cor do PDF (#f3eee2)
    bg = OxmlElement("w:background")
    bg.set(qn("w:color"), "F3EEE2")
    doc.element.insert(0, bg)
    try:
        disp = OxmlElement("w:displayBackgroundShape")
        doc.settings.element.append(disp)
    except Exception:
        pass

    ctx = _build_context(proposal, session)
    price_str = ctx["{{total_price}}"]
    days_str = ctx["{{total_days}}"]
    client_name = ctx["{{client_name}}"]
    analyst_name = proposal.get("analyst_name") or "—"

    created_at_raw = proposal.get("created_at")
    try:
        today = (
            datetime.fromisoformat(str(created_at_raw).replace("Z", "+00:00")).strftime("%d/%m/%Y")
            if created_at_raw
            else datetime.now(timezone.utc).strftime("%d/%m/%Y")
        )
    except Exception:
        today = datetime.now(timezone.utc).strftime("%d/%m/%Y")

    raw_meta = session.get("proposal_metadata") or ""
    content = _strip_metadata_block(raw_meta) if raw_meta else ""
    sections_list = _parse_sections(content) if content else []

    C_RED   = RGBColor(0xC0, 0x39, 0x2B)
    C_DARK  = RGBColor(0x1A, 0x1A, 0x1A)
    C_GRAY  = RGBColor(0x55, 0x55, 0x55)
    C_LGRAY = RGBColor(0x99, 0x99, 0x99)
    C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    C_MUTED = RGBColor(0xAA, 0xA5, 0x9C)  # texto secundário no fundo escuro

    def _r(run, size: int, color: RGBColor, bold=False, font="Arial"):
        run.font.size = Pt(size)
        run.font.name = font
        run.font.color.rgb = color
        if bold:
            run.bold = True

    def _add_runs(para, text: str, size=10, color: RGBColor = None, font="Arial"):
        c = color or C_DARK
        for m in re.finditer(r"\*\*([^*]+)\*\*|\*([^*]+)\*|([^*\n]+)", text):
            if m.group(1):
                run = para.add_run(m.group(1))
                _r(run, size, c, bold=True, font=font)
            elif m.group(2):
                run = para.add_run(m.group(2))
                _r(run, size, c, font=font)
                run.italic = True
            elif m.group(3):
                run = para.add_run(m.group(3))
                _r(run, size, c, font=font)

    def _sp(para, before=0, after=6, line=14):
        pf = para.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = Pt(line)

    def _add_hr(para, color="D9D2BF", sz=4):
        pPr = para._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        el = OxmlElement("w:bottom")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "1")
        el.set(qn("w:color"), color)
        pBdr.append(el)
        pPr.append(pBdr)

    def _add_cell_top_border(cell, color="1A1A1A", sz=8):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBorders = OxmlElement("w:tcBorders")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), str(sz))
        top.set(qn("w:space"), "0")
        top.set(qn("w:color"), color)
        tcBorders.append(top)
        tcPr.append(tcBorders)

    def _set_cell_bg(cell, fill: str):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    def _clear_cell_borders(cell):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBorders = OxmlElement("w:tcBorders")
        for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "none")
            tcBorders.append(el)
        tcPr.append(tcBorders)

    def add_label(heading: str, num: int):
        p = doc.add_paragraph()
        _sp(p, before=0, after=10)
        run = p.add_run(f"{str(num).zfill(2)} · {heading.upper()}")
        _r(run, 8, C_RED, bold=True, font="Courier New")

    def add_body(text: str):
        p = doc.add_paragraph()
        _sp(p, after=6)
        _add_runs(p, text)

    def add_bullet(text: str):
        p = doc.add_paragraph(style="List Bullet")
        _sp(p, after=3)
        _add_runs(p, text, size=10)

    def add_numbered(text: str):
        p = doc.add_paragraph(style="List Number")
        _sp(p, after=3)
        _add_runs(p, text, size=10)

    def add_sub(text: str):
        p = doc.add_paragraph()
        _sp(p, before=8, after=3)
        run = p.add_run(text)
        _r(run, 10, C_DARK, bold=True)

    def render_body(body: str):
        for line in body.split("\n"):
            t = line.strip()
            if not t:
                continue
            if re.match(r"^[-*_]{3,}$", t):
                p = doc.add_paragraph()
                _sp(p, before=4, after=4)
                _add_hr(p)
                continue
            hm = re.match(r"^(#{1,4})\s+(.+)$", t)
            if hm:
                add_sub(hm.group(2))
                continue
            bm = re.match(r"^[-•*]\s+(.+)$", t)
            if bm:
                if _is_empty_value(bm.group(1)):
                    continue
                add_bullet(bm.group(1))
                continue
            om = re.match(r"^\d+[.)]\s+(.+)$", t)
            if om:
                add_numbered(om.group(1))
                continue
            add_body(t)

    # ── Capa ──────────────────────────────────────────────────────────
    p = doc.add_paragraph()
    _sp(p, before=0, after=6)
    _r(p.add_run("Archi"), 14, C_DARK, bold=True)

    p = doc.add_paragraph()
    _sp(p, after=20)
    _r(p.add_run(f"PROPOSTA COMERCIAL · {today}"), 8, C_LGRAY, font="Courier New")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(10)
    # line_spacing não definido: Word usa automático (correto para fonte 24pt)
    title_text = (
        f"Proposta para {client_name}."
        if client_name and client_name != "—"
        else "Proposta Técnica."
    )
    _r(p.add_run(title_text), 24, C_DARK, bold=True, font="Georgia")

    p = doc.add_paragraph()
    _sp(p, after=28)
    sub_text = (
        f"Proposta comercial elaborada com escopo, prazo e investimento para {client_name}."
        if client_name and client_name != "—"
        else "Proposta técnica com escopo, prazo e investimento detalhados."
    )
    _r(p.add_run(sub_text), 11, C_GRAY)

    # Bloco escuro da capa (espelho do rodapé escuro do PDF)
    tc = doc.add_table(rows=1, cols=2)
    tc.style = "Table Grid"
    for cell in tc.rows[0].cells:
        _set_cell_bg(cell, "1A1A1A")
        _clear_cell_borders(cell)

    cl = tc.cell(0, 0)
    _sp(cl.paragraphs[0], before=10, after=10)
    _r(cl.paragraphs[0].add_run("PREPARADO PARA\n"), 7, C_MUTED, font="Courier New")
    _r(cl.paragraphs[0].add_run(client_name), 13, C_WHITE, bold=True)

    cr = tc.cell(0, 1)
    _sp(cr.paragraphs[0], before=10, after=10)
    _r(cr.paragraphs[0].add_run("INVESTIMENTO\n"), 7, C_MUTED, font="Courier New")
    _r(cr.paragraphs[0].add_run(price_str), 14, C_WHITE, bold=True, font="Courier New")
    if days_str and days_str != "A definir":
        _r(cr.paragraphs[0].add_run(f"\n{days_str} dias úteis"), 9, C_MUTED)

    doc.add_page_break()

    # ── Seções de conteúdo (fluxo contínuo, sem page break forçado) ───
    for i, sec in enumerate(sections_list):
        if i > 0:
            p = doc.add_paragraph()
            _sp(p, before=10, after=8)
            _add_hr(p, color="D9D2BF", sz=6)
        add_label(sec["heading"], i + 1)
        render_body(_substitute_vars(sec["body"], proposal))

    # ── Assinaturas ───────────────────────────────────────────────────
    doc.add_page_break()
    add_label("ASSINATURAS", len(sections_list) + 1)

    p = doc.add_paragraph()
    _sp(p, after=24)
    _r(
        p.add_run(
            "Ao assinar abaixo, as partes declaram estar de acordo com o escopo, "
            "prazo e investimento descritos nesta proposta."
        ),
        10, C_DARK,
    )

    p = doc.add_paragraph()
    _sp(p, after=40)
    _r(p.add_run("Data: ____/____/______"), 10, C_GRAY)

    sig = doc.add_table(rows=2, cols=3)
    for cell in sig.rows[0].cells:
        cell.paragraphs[0].add_run("\n\n\n").font.size = Pt(10)

    _add_cell_top_border(sig.cell(1, 0))
    _add_cell_top_border(sig.cell(1, 2))

    p_l = sig.cell(1, 0).paragraphs[0]
    _r(p_l.add_run(client_name), 10, C_DARK, bold=True)
    _r(sig.cell(1, 0).add_paragraph().add_run("CONTRATANTE"), 8, C_LGRAY, font="Courier New")

    p_r = sig.cell(1, 2).paragraphs[0]
    _r(p_r.add_run(analyst_name), 10, C_DARK, bold=True)
    _r(sig.cell(1, 2).add_paragraph().add_run("CONTRATADO"), 8, C_LGRAY, font="Courier New")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF generation — content sourced from proposal_metadata (agent output)
# ---------------------------------------------------------------------------

def _he(s: object) -> str:
    return _html.escape(str(s)) if s else ""


def _strip_metadata_block(raw: str) -> str:
    return re.sub(r"\[PROPOSAL_METADATA\][\s\S]*?\[\/PROPOSAL_METADATA\]", "", raw, flags=re.IGNORECASE).strip()


def _substitute_vars(text: str, proposal: dict) -> str:
    from app.config import settings  # local import to avoid circular dep
    price = proposal.get("total_price")
    if price:
        price_str = f"R$ {price:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    else:
        price_str = "—"
    days = proposal.get("total_days")
    validity = str(proposal.get("validity_days") or 30)
    return (
        text
        .replace("{{NOME_CLIENTE}}", proposal.get("client_name") or "—")
        .replace("{{NOME_ANALISTA}}", proposal.get("analyst_name") or "—")
        .replace("{{PRECO_TOTAL}}", price_str)
        .replace("{{PRAZO_DIAS}}", f"{days} dias úteis" if days else "—")
        .replace("{{EMAIL_ANALISTA}}", settings.company_email or "—")
        .replace("{{TELEFONE_ANALISTA}}", settings.company_phone or "—")
        .replace("{{SITE_EMPRESA}}", settings.company_website or "—")
        .replace("[___]", validity)
    )


def _parse_sections(content: str) -> list[dict]:
    sections: list[dict] = []
    parts = re.split(r"(?=^###\s)", content, flags=re.MULTILINE)
    for part in parts:
        stripped = part.strip()
        if not stripped:
            continue
        m = re.match(r"^###\s+(?:\d+\.?\s+)?(.+?)(?:\n|$)", stripped)
        if m:
            heading = m.group(1).strip()
            body = stripped[stripped.index("\n") + 1:].strip() if "\n" in stripped else ""
            sections.append({"heading": heading, "body": body})
    return sections


def _md_inline(text: str) -> str:
    escaped = _he(text)
    escaped = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", escaped)
    escaped = re.sub(r"\*([^*]+)\*", r"<em>\1</em>", escaped)
    return escaped


def _is_empty_value(item_text: str) -> bool:
    """Return True if a bullet item has an empty value (resolved to '—')."""
    clean = re.sub(r"\*+", "", item_text).strip()
    return bool(re.search(r":\s*—\s*$", clean)) or clean == "—"


def _render_section_body(body: str) -> str:
    lines = body.split("\n")
    html_parts: list[str] = []
    bullet_items: list[str] = []
    ordered_items: list[str] = []

    def flush_ul():
        if bullet_items:
            items = "".join(f'<li style="font-size:10pt;line-height:1.65;color:#333;margin-bottom:3pt;">{_md_inline(i)}</li>' for i in bullet_items)
            html_parts.append(f'<ul style="padding-left:14pt;margin:3pt 0 8pt;">{items}</ul>')
            bullet_items.clear()

    def flush_ol():
        if ordered_items:
            items = "".join(f'<li style="font-size:10pt;line-height:1.65;color:#333;margin-bottom:3pt;">{_md_inline(i)}</li>' for i in ordered_items)
            html_parts.append(f'<ol style="padding-left:16pt;margin:3pt 0 8pt;">{items}</ol>')
            ordered_items.clear()

    for line in lines:
        t = line.strip()
        if not t:
            flush_ul(); flush_ol()
            continue
        if re.match(r"^[-*_]{3,}$", t):
            flush_ul(); flush_ol()
            html_parts.append('<hr style="border:none;border-top:1px solid #ddd;margin:12pt 0;">')
            continue
        hm = re.match(r"^(#{1,4})\s+(.+)$", t)
        if hm:
            flush_ul(); flush_ol()
            fs = "12pt" if len(hm.group(1)) <= 2 else "10.5pt"
            html_parts.append(f'<h4 style="font-size:{fs};font-weight:600;margin:10pt 0 3pt;color:#1a1a1a;">{_md_inline(hm.group(2))}</h4>')
            continue
        bm = re.match(r"^[-•*]\s+(.+)$", t)
        if bm:
            if _is_empty_value(bm.group(1)):
                continue
            flush_ol()
            bullet_items.append(bm.group(1))
            continue
        om = re.match(r"^\d+[.)]\s+(.+)$", t)
        if om:
            flush_ul()
            ordered_items.append(om.group(1))
            continue
        flush_ul(); flush_ol()
        html_parts.append(f'<p style="font-size:10pt;line-height:1.65;color:#333;margin:0 0 7pt;">{_md_inline(t)}</p>')

    flush_ul(); flush_ol()
    return "\n".join(html_parts)


def _generate_pdf_bytes(proposal: dict, session: dict) -> bytes:
    ctx = _build_context(proposal, session)
    price_str = ctx["{{total_price}}"]
    days_str = ctx["{{total_days}}"]
    client_name = ctx["{{client_name}}"]
    analyst_name = proposal.get("analyst_name") or "—"

    created_at_raw = proposal.get("created_at")
    try:
        today = datetime.fromisoformat(
            str(created_at_raw).replace("Z", "+00:00")
        ).strftime("%d/%m/%Y") if created_at_raw else datetime.now(timezone.utc).strftime("%d/%m/%Y")
    except Exception:
        today = datetime.now(timezone.utc).strftime("%d/%m/%Y")

    # --- Content from proposal_metadata (final agent output) ---
    raw_meta = session.get("proposal_metadata") or ""
    content = _strip_metadata_block(raw_meta) if raw_meta else ""
    sections = _parse_sections(content) if content else []

    # Extract phase count from [PROPOSAL_METADATA] block
    phase_count_m = re.search(r"numero_de_fases:\s*(\d+)", raw_meta, re.IGNORECASE)
    phase_count = int(phase_count_m.group(1)) if phase_count_m else 0
    phase_info = f" · {phase_count} fases" if phase_count > 0 else ""

    title = (
        f"Proposta para {client_name}."
        if client_name and client_name != "—"
        else "Proposta Técnica."
    )
    subtitle = (
        f"Proposta comercial elaborada com escopo, prazo e investimento para {client_name}."
        if client_name and client_name != "—"
        else "Proposta técnica com escopo, prazo e investimento detalhados."
    )

    # Build section HTML blocks — flow continuously, WeasyPrint paginates automatically
    section_blocks: list[str] = []
    for i, sec in enumerate(sections):
        top_margin = "margin-top:28pt;" if i > 0 else ""
        body = _substitute_vars(sec["body"], proposal)
        body_html = _render_section_body(body)
        section_blocks.append(f"""
<div style="{top_margin}padding:0;">
  <div style="font-size:8pt;font-family:Courier New,monospace;letter-spacing:0.12em;text-transform:uppercase;color:#C0392B;margin-bottom:10pt;">
    {str(i + 1).zfill(2)} · {_he(sec['heading'])}
  </div>
  {body_html}
</div>""")

    sig_num = str(len(sections) + 1).zfill(2)
    signatures_html = f"""
<div style="page-break-before:always;padding:0;">
  <div style="font-size:8pt;font-family:Courier New,monospace;letter-spacing:0.12em;text-transform:uppercase;color:#C0392B;margin-bottom:16pt;">
    {sig_num} · ASSINATURAS
  </div>
  <p style="font-size:10pt;line-height:1.65;color:#333;margin:0 0 32pt;">
    Ao assinar abaixo, as partes declaram estar de acordo com o escopo, prazo e investimento descritos nesta proposta.
  </p>
  <p style="font-size:9pt;color:#888;font-family:Courier New,monospace;margin:0 0 52pt;">Data: ____/____/______</p>
  <table style="width:100%;border-collapse:collapse;">
    <tr>
      <td style="width:44%;vertical-align:top;border-top:1pt solid #1a1a1a;padding-top:8pt;padding-right:20pt;">
        <div style="font-size:10pt;font-weight:600;color:#1a1a1a;">{_he(client_name)}</div>
        <div style="font-size:8pt;color:#999;font-family:Courier New,monospace;letter-spacing:0.08em;text-transform:uppercase;margin-top:4pt;">Contratante</div>
      </td>
      <td style="width:12%;"></td>
      <td style="width:44%;vertical-align:top;border-top:1pt solid #1a1a1a;padding-top:8pt;padding-left:20pt;">
        <div style="font-size:10pt;font-weight:600;color:#1a1a1a;">{_he(analyst_name)}</div>
        <div style="font-size:8pt;color:#999;font-family:Courier New,monospace;letter-spacing:0.08em;text-transform:uppercase;margin-top:4pt;">Contratado</div>
      </td>
    </tr>
  </table>
</div>"""

    sections_html = ("\n".join(section_blocks) + signatures_html) if section_blocks else (
        "<p style='font-size:10pt;color:#888;font-style:italic;'>Conteúdo da proposta não disponível.</p>"
        + signatures_html
    )

    html_content = f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="UTF-8">
<style>
  @page {{
    size: A4;
    margin: 16mm 18mm 14mm;
    background: #f3eee2;
    @bottom-left {{ content: "Archi · Proposta Comercial"; font-family: "Courier New", monospace; font-size: 7pt; color: #999; }}
    @bottom-right {{ content: counter(page) " / " counter(pages); font-family: "Courier New", monospace; font-size: 7pt; color: #999; }}
  }}
  @page :first {{ margin: 0; background: #f3eee2; @bottom-left {{ content: none; }} @bottom-right {{ content: none; }} }}
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ font-family: Arial, Helvetica, sans-serif; color: #1a1a1a; background: #f3eee2; }}
</style>
</head>
<body>

<!-- COVER (page 1 — full bleed via @page :first) -->
<div style="width:210mm;height:297mm;position:relative;overflow:hidden;background:#f3eee2;page-break-after:always;">
  <div style="padding:28mm 18mm 0;">
    <div style="font-size:13pt;font-weight:700;letter-spacing:-0.01em;margin-bottom:10mm;">Archi</div>
    <div style="font-size:8pt;font-family:Courier New,monospace;letter-spacing:0.1em;text-transform:uppercase;color:#888;margin-bottom:14mm;">Proposta Comercial · {_he(today)}</div>
    <h1 style="font-family:Georgia,serif;font-size:26pt;font-weight:700;line-height:1.1;letter-spacing:-0.03em;color:#1a1a1a;margin-bottom:8mm;max-width:130mm;">{_he(title)}</h1>
    <p style="font-size:11pt;line-height:1.6;color:#555;max-width:120mm;">{_he(subtitle)}</p>
  </div>
  <div style="position:absolute;bottom:0;left:0;right:0;background:#1a1a1a;padding:10mm 18mm;">
    <table style="width:100%;border-collapse:collapse;"><tr>
      <td style="vertical-align:bottom;padding:0;">
        <div style="font-size:7pt;font-family:Courier New,monospace;letter-spacing:0.1em;text-transform:uppercase;color:rgba(255,255,255,0.45);margin-bottom:2mm;">Preparado para</div>
        <div style="font-size:13pt;color:white;font-weight:600;">{_he(client_name)}</div>
      </td>
      <td style="vertical-align:bottom;text-align:right;padding:0;">
        <div style="font-size:7pt;font-family:Courier New,monospace;letter-spacing:0.1em;text-transform:uppercase;color:rgba(255,255,255,0.45);margin-bottom:2mm;">Investimento</div>
        <div style="font-size:16pt;color:white;font-weight:700;font-family:Courier New,monospace;">{_he(price_str)}</div>
        <div style="font-size:9pt;color:rgba(255,255,255,0.65);margin-top:1mm;">{_he(days_str)} dias úteis{_he(phase_info)}</div>
      </td>
    </tr></table>
  </div>
</div>

<!-- PROPOSAL CONTENT (sections flow across as many pages as needed) -->
{sections_html}

</body>
</html>"""

    return HTML(string=html_content).write_pdf()


async def generate_docs(
    proposal_id: str,
    fmt: str | None = None,
    overrides: dict | None = None,
) -> dict:
    """Generate documents for a proposal.

    fmt: 'docx', 'pdf', or None for both.
    overrides: values from the current UI state that override the DB values.
    """
    sb = get_supabase()

    proposal_res = await sb.table("proposals").select("*").eq("id", proposal_id).execute()
    if not proposal_res.data:
        raise ValueError(f"Proposta {proposal_id} não encontrada.")
    proposal = proposal_res.data[0]

    # Apply UI state overrides so PDF reflects what admin sees on screen
    if overrides:
        proposal = {**proposal, **overrides}

    session_res = await sb.table("sessions").select("*").eq("id", proposal["session_id"]).execute()
    session = session_res.data[0] if session_res.data else {}

    session_id = proposal_res.data[0]["session_id"]
    result: dict = {}

    if fmt in (None, "docx"):
        docx_path = f"{session_id}/proposal.docx"
        docx_bytes = _generate_docx_bytes(proposal, session)
        try:
            await sb.storage.from_(BUCKET).remove([docx_path])
        except Exception:
            pass
        await sb.storage.from_(BUCKET).upload(
            docx_path,
            docx_bytes,
            {"content-type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
        )
        result["docx_url"] = await sb.storage.from_(BUCKET).get_public_url(docx_path)

    if fmt in (None, "pdf"):
        pdf_path = f"{session_id}/proposal.pdf"
        pdf_bytes = _generate_pdf_bytes(proposal, session)
        try:
            await sb.storage.from_(BUCKET).remove([pdf_path])
        except Exception:
            pass
        await sb.storage.from_(BUCKET).upload(
            pdf_path,
            pdf_bytes,
            {"content-type": "application/pdf"},
        )
        result["pdf_url"] = await sb.storage.from_(BUCKET).get_public_url(pdf_path)

    logger.info(f"Documentos gerados para proposta {proposal_id} (fmt={fmt}): {result}")
    return result
